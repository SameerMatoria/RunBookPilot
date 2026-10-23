"""Turn an uploaded runbook file (PDF, Word, Markdown, text) into plain text."""
import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(io.BytesIO(content))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix in {".md", ".txt", ""}:
        return content.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}. Use PDF, DOCX, MD or TXT.")
