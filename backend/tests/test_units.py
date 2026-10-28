"""Unit tests for the lint gate, sandbox log parsing, documents, queue and metrics."""
import io

from docx import Document

from app import jobqueue, metrics, store
from app.documents import extract_text
from app.lint import check_blocked, lint_script
from app.sandbox import _parse_steps
from conftest import SAFE_SCRIPT, upload


# ---------- lint gate ----------

def test_blocked_commands_are_caught():
    for line in ["rm -rf /", "mkfs.ext4 /dev/sdb1", "dd if=/dev/zero of=/dev/sda", "iptables -F",
                 "shutdown -h now", 'PASSWORD="hunter22"', "chmod -R 777 /etc"]:
        assert check_blocked(line), line


def test_safe_commands_pass():
    for line in ["rm -f /var/log/nginx/old.gz", "systemctl restart nginx", 'password="$DB_PASSWORD"',
                 "# rm -rf / in a comment is fine"]:
        assert not check_blocked(line), line


def test_lint_script_passes_good_script_and_flags_syntax_error():
    assert lint_script(SAFE_SCRIPT)["passed"] is True
    broken = lint_script("#!/usr/bin/env bash\nif [ 1 -eq 1 ]; then\n  echo hi\n")
    assert broken["passed"] is False


# ---------- sandbox log parsing ----------

def test_parse_steps_takes_final_status_and_marks_crashes():
    out = "noise\nRBP_STEP|1|START|a\nRBP_STEP|1|OK|b\nRBP_STEP|2|START|c\nRBP_STEP|3|SKIP|d"
    steps = {s["order"]: s["status"] for s in _parse_steps(out)}
    assert steps == {1: "OK", 2: "FAIL", 3: "SKIP"}


# ---------- documents ----------

def test_extract_markdown_and_docx():
    assert "Restart" in extract_text("rb.md", b"# Restart nginx")
    doc = Document()
    doc.add_paragraph("Step 1: check disk")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Step 2"
    table.cell(0, 1).text = "restart"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("rb.docx", buf.getvalue())
    assert "check disk" in text and "Step 2 | restart" in text


# ---------- queue ----------

def test_queue_claims_each_job_once():
    rid = store.create("rb.md", "text")
    first = jobqueue.enqueue("parse", rid)
    jobqueue.enqueue("build", rid, feedback="x")
    assert jobqueue.has_active_job(rid)

    a = jobqueue.claim("w1")
    b = jobqueue.claim("w2")
    assert a["id"] == first and b["kind"] == "build" and b["payload"] == {"feedback": "x"}
    assert jobqueue.claim("w3") is None
    jobqueue.finish(a["id"])
    jobqueue.finish(b["id"], error="boom")
    assert not jobqueue.has_active_job(rid)


# ---------- metrics ----------

def test_metrics_hours_saved(client, fakes):
    upload(client)
    m = metrics.compute()
    assert m["totals"]["processed"] == 1
    assert m["totals"]["converted"] == 1
    assert m["totals"]["first_attempt"] == 1
    row = m["runbooks"][0]
    # (15 manual - 2 automated) min x 30 runs/month x 12 / 60
    assert row["hours_saved_per_year"] == 78.0
    assert row["speedup"] == 7.5


# ---------- worker heartbeats ----------

def test_sandbox_status_comes_from_worker_heartbeats():
    from app import heartbeat

    assert heartbeat.status() == {"ready": False, "workers": 0, "detail": heartbeat.status()["detail"]}
    heartbeat.beat("w1", False, "Docker engine is not running.")
    assert heartbeat.status()["ready"] is False and heartbeat.status()["workers"] == 1
    heartbeat.beat("w2", True, "Docker 29.9")
    assert heartbeat.status() == {"ready": True, "workers": 2, "detail": "Sandbox ready"}
    heartbeat.remove("w1")
    heartbeat.remove("w2")
    assert heartbeat.status()["workers"] == 0
